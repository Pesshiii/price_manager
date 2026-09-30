from decimal import Decimal
from io import BytesIO
from unittest import mock

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from main_product_manager.models import MainProduct
from pim_api import FlatEntityList
from supplier_manager.models import Supplier

from product.models import Brand, Product, ProductExport
from product.satu_export import (
    BASE_TITLES, CHARACTERISTIC_TITLES, GROUP_TITLES, GROUPS_SHEET, MIN_CHARACTERISTICS,
    PRICE_TITLES, PRODUCTS_SHEET, build_satu_export,
)

ATTRIBUTES = [
    {'id': 'a-len', 'name': 'Длина', 'sortOrder': 2},
    {'id': 'a-mat', 'name': 'Материал', 'sortOrder': 1},
    {'id': 'a-cord', 'name': 'Сетевой', 'sortOrder': 3},
]

FLAT = {
    'id': 'pim-1',
    'countryOfOriginName': 'Германия',
    'ean': '4600000000001',
    'mpn': None,
    'attributesDefs': {
        'length': {'attributeId': 'a-len', 'type': 'float'},
        'lengthUnit': {'attributeId': 'a-len', 'type': 'link'},
        'material': {'attributeId': 'a-mat', 'type': 'varchar'},
        'corded': {'attributeId': 'a-cord', 'type': 'bool'},
        'empty': {'attributeId': 'a-x', 'type': 'varchar'},
        'setProducts': {'attributeId': 'a-set', 'type': 'linkMultiple'},
    },
    'length': 1.5,
    'lengthUnit': 'u-1',
    'lengthUnitData': {'symbol': 'м'},
    'material': 'Сталь',
    'corded': False,
    'empty': ' ',
    'setProducts': ['x'],
}


def fake_pim(query, timeout=None):
    if query.name == 'Attribute':
        return {'total': len(ATTRIBUTES), 'list': ATTRIBUTES if not query.offset else []}
    assert isinstance(query, FlatEntityList)
    ids = query.where[0].value
    return {'total': 1, 'list': [FLAT] if 'pim-1' in ids else []}


class SatuExportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(username='admin', password='pw')
        self.a = Supplier.objects.create(name='А', price_priority=1, stock_priority=1,
                                         delivery_days_navailable=7)
        self.b = Supplier.objects.create(name='Б', price_priority=2, stock_priority=2,
                                         delivery_days_navailable=3)
        brand = Brand.objects.create(pim_id='b-1', name='Зубр')
        self.product = Product.objects.create(
            pim_id='p-1', number='SKU-1', name='Молоток', brand=brand,
            raw_data={'id': 'pim-1', 'longDescription': '<p>Кованый</p>',
                      'description': 'коротко', 'tag': ['молоток', 'инструмент']})
        MainProduct.objects.create(product=self.product, supplier=self.a, article='A-1',
                                   name='Молоток А', m_price=Decimal('1500'), stock=4)
        self.bare = Product.objects.create(pim_id='p-2', number='SKU-2', name='Без PIM')
        MainProduct.objects.create(product=self.bare, supplier=self.a, article='A-2',
                                   name='Без PIM А', stock=0)
        MainProduct.objects.create(product=self.bare, supplier=self.b, article='B-2',
                                   name='Без PIM Б', stock=0)

    def export(self):
        with mock.patch('product.pim_client.site', **{'get.side_effect': fake_pim}):
            export = build_satu_export(self.user.pk)
        with export.file.open('rb') as file:
            workbook = load_workbook(BytesIO(file.read()), read_only=True)
        sheets = {}
        for sheet in workbook.worksheets:
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            # read_only обрезает пустые ячейки в конце строки.
            sheets[sheet.title] = [row + [None] * (len(rows[0]) - len(row)) for row in rows]
        return export, sheets

    def test_layout_and_values(self):
        export, sheets = self.export()
        self.assertEqual(list(sheets), [PRODUCTS_SHEET, GROUPS_SHEET])
        self.assertEqual(sheets[GROUPS_SHEET], [GROUP_TITLES])
        self.assertEqual(export.rows_count, 2)
        self.assertTrue(export.file.name.endswith('.xlsx'))

        header, *rows = sheets[PRODUCTS_SHEET]
        self.assertEqual(header, BASE_TITLES + CHARACTERISTIC_TITLES * MIN_CHARACTERISTICS
                         + PRICE_TITLES)
        by_code = {row[0]: dict(zip(BASE_TITLES, row)) for row in rows}
        row = by_code['SKU-1']
        self.assertEqual(row['Название_позиции'], 'Молоток')
        self.assertEqual(row['Описание'], '<p>Кованый</p>')
        self.assertEqual(row['Поисковые_запросы'], 'молоток, инструмент')
        self.assertEqual(row['Производитель'], 'Зубр')
        self.assertEqual(row['Страна_производитель'], 'Германия')
        self.assertEqual(row['Код_маркировки_(GTIN)'], '4600000000001')
        self.assertIsNone(row['Номер_устройства_(MPN)'])
        self.assertEqual(row['Наличие'], '+')
        self.assertEqual(row['Количество'], 4)
        # Цены Satu пустые — все цены ГП в своих колонках в конце.
        self.assertIsNone(row['Цена'])
        full = next(r for r in rows if r[0] == 'SKU-1')
        prices = dict(zip(PRICE_TITLES, full[-len(PRICE_TITLES):]))
        self.assertEqual(prices['Цена ИМ'], 1500)

        # Характеристики — по порядку атрибутов PIM, без пустых, единиц-ссылок и связей.
        start = len(BASE_TITLES)
        self.assertEqual(full[start:start + 9], [
            'Материал', None, 'Сталь', 'Длина', 'м', '1,5', 'Сетевой', None, 'Нет'])
        self.assertIsNone(full[start + 9])

    def test_without_stock_availability_is_fastest_delivery(self):
        _, sheets = self.export()
        row = next(r for r in sheets[PRODUCTS_SHEET][1:] if r[0] == 'SKU-2')
        row = dict(zip(BASE_TITLES, row))
        self.assertEqual(row['Наличие'], 3)
        self.assertEqual(row['Количество'], 0)
        self.assertIsNone(row['Страна_производитель'])

    def test_no_stock_and_no_delivery_is_unavailable(self):
        Supplier.objects.update(delivery_days_navailable=None)
        _, sheets = self.export()
        row = next(r for r in sheets[PRODUCTS_SHEET][1:] if r[0] == 'SKU-2')
        self.assertEqual(dict(zip(BASE_TITLES, row))['Наличие'], '-')

    def test_control_characters_are_stripped(self):
        # Такие символы приходят из текста, скопированного из Word; openpyxl
        # на них падает — и роняет выгрузку в самом конце, после всех запросов к PIM.
        self.product.raw_data['longDescription'] = '<p>Кованыйбоёк</p>'
        self.product.save()
        flat = dict(FLAT, material='Сталь')
        with mock.patch('product.pim_client.site', **{'get.side_effect': fake_pim}),                 mock.patch.dict(FLAT, flat):
            export = build_satu_export(self.user.pk)
        with export.file.open('rb') as file:
            rows = list(load_workbook(BytesIO(file.read()), read_only=True)
                        [PRODUCTS_SHEET].iter_rows(values_only=True))
        row = next(r for r in rows if r[0] == 'SKU-1')
        self.assertEqual(dict(zip(BASE_TITLES, row))['Описание'], '<p>Кованыйбоёк</p>')
        self.assertIn('Сталь', row)

    def test_wider_than_satu_when_product_has_more_characteristics(self):
        many = dict(FLAT, attributesDefs={f'c{i}': {'attributeId': f'x{i}', 'type': 'int'}
                                          for i in range(MIN_CHARACTERISTICS + 2)})
        many.update({f'c{i}': i for i in range(MIN_CHARACTERISTICS + 2)})
        with mock.patch('product.satu_export.fetch_flat_products',
                        return_value={'pim-1': many}), \
                mock.patch('product.satu_export.fetch_attributes', return_value={}):
            export = build_satu_export(self.user.pk)
        with export.file.open('rb') as file:
            header = next(load_workbook(BytesIO(file.read()), read_only=True)
                          [PRODUCTS_SHEET].iter_rows(values_only=True))
        self.assertEqual(header.count('Название_Характеристики'), MIN_CHARACTERISTICS + 2)

    def test_admin_button_posts_and_queues_task(self):
        self.client.force_login(self.user)
        changelist = reverse('admin:product_product_changelist')
        url = reverse('admin:product_product_export_satu')
        self.assertContains(self.client.get(changelist), url)
        self.assertEqual(self.client.get(url).status_code, 405)
        with mock.patch('product.admin.export_products_satu_task.delay') as delay:
            response = self.client.post(url)
        self.assertRedirects(response, changelist)
        delay.assert_called_once_with(user_id=self.user.pk)

    def test_admin_export_needs_view_permission(self):
        staff = User.objects.create_user(username='staff', password='pw', is_staff=True)
        self.client.force_login(staff)
        with mock.patch('product.admin.export_products_satu_task.delay') as delay:
            response = self.client.post(reverse('admin:product_product_export_satu'))
        self.assertEqual(response.status_code, 403)
        delay.assert_not_called()

    def test_task_notifies_with_download_link(self):
        from core.models import PersistentNotification
        from product.tasks import export_products_satu_task

        with mock.patch('product.pim_client.site', **{'get.side_effect': fake_pim}):
            export_products_satu_task(user_id=self.user.pk)
        export = ProductExport.objects.get(user=self.user)
        notification = PersistentNotification.objects.get(user=self.user)
        self.assertEqual(notification.link,
                         reverse('product-export-download', kwargs={'pk': export.pk}))
