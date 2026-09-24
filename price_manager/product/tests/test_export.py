import csv
from decimal import Decimal
from io import BytesIO, StringIO
from unittest import mock

from django.contrib.auth.models import User
from django.core.files.base import ContentFile
from django.http import QueryDict
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from main_product_manager.models import MainProduct
from supplier_manager.models import Supplier

from product.columns import COLUMN_LABELS
from product.export import (
    FULL_EXPORT_COLUMNS, MAIN_PRICE_COLUMNS, ProductExporter, build_full_csv_export,
    main_value, ordered_product_pks, ranked_suppliers, sheet_names, supplier_levels,
)
from product.models import Brand, Category, Product, ProductExport


def read_sheets(content):
    """{имя листа: [строка, …]} в порядке листов книги.

    read_only обрезает пустые ячейки в конце строки — строки дополняются до
    ширины заголовка, иначе пустой остаток был бы неотличим от пропавшей колонки.
    """
    workbook = load_workbook(BytesIO(content), read_only=True)
    sheets = {}
    for sheet in workbook.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        width = len(rows[0]) if rows else 0
        sheets[sheet.title] = [row + [None] * (width - len(row)) for row in rows]
    return sheets


def as_dicts(rows):
    header, *body = rows
    return [dict(zip(header, row)) for row in body]


class MainValueTests(TestCase):
    def test_single_supplier_levels_first_non_zero_wins(self):
        self.assertEqual(main_value([[Decimal('0')], [None], [Decimal('7')]], min), Decimal('7'))

    def test_min_or_max_within_level(self):
        levels = [[Decimal('5'), Decimal('3'), Decimal('9')], [Decimal('1')]]
        self.assertEqual(main_value(levels, min), Decimal('3'))
        self.assertEqual(main_value(levels, max), Decimal('9'))

    def test_zero_and_empty_skipped_within_level(self):
        # Ноль не «самая низкая цена» — это «цены нет».
        self.assertEqual(main_value([[0, None, Decimal('4')]], min), Decimal('4'))

    def test_level_without_values_falls_through(self):
        self.assertEqual(main_value([[0, None], [Decimal('2'), Decimal('8')]], max), Decimal('8'))

    def test_all_zero_or_empty_gives_zero(self):
        self.assertEqual(main_value([[None, 0], [None]], min), 0)

    def test_all_empty_gives_empty(self):
        self.assertIsNone(main_value([[None], [None, None]], min))

    def test_no_suppliers_gives_empty(self):
        self.assertIsNone(main_value([], min))


class RankedSuppliersTests(TestCase):
    def test_lower_priority_first_unranked_and_no_supplier_last(self):
        unranked = Supplier.objects.create(name='А-без приоритета')
        second = Supplier.objects.create(name='Б', price_priority=2)
        first = Supplier.objects.create(name='В', price_priority=1)
        ranked = ranked_suppliers({unranked.pk, second.pk, first.pk, None}, 'price_priority')
        self.assertEqual([name for _, name in ranked], ['В', 'Б', 'А-без приоритета', 'Без поставщика'])


class SupplierLevelsTests(TestCase):
    def test_shared_numbers_group_unranked_share_one_level_no_supplier_last(self):
        a = Supplier.objects.create(name='А', price_priority=5)
        b = Supplier.objects.create(name='Б', price_priority=1)
        c = Supplier.objects.create(name='В', price_priority=5)
        d = Supplier.objects.create(name='Г')
        e = Supplier.objects.create(name='Д')
        levels = supplier_levels({a.pk, b.pk, c.pk, d.pk, e.pk, None}, 'price_priority')
        self.assertEqual([sorted(level, key=str) for level in levels],
                         [[b.pk], sorted([a.pk, c.pk], key=str),
                          sorted([d.pk, e.pk], key=str), [None]])

    def test_no_suppliers(self):
        self.assertEqual(supplier_levels(set(), 'stock_priority'), [])


class SheetNamesTests(TestCase):
    def test_forbidden_chars_length_and_case_insensitive_duplicates(self):
        long = 'Очень длинное название поставщика ТОО'
        names = sheet_names(['ИП «Рога/Копыта» [опт]', long, long.upper(), 'товары', ''])
        self.assertEqual(names[0], 'ИП «Рога Копыта»  опт')
        self.assertEqual(names[1], long[:31])
        self.assertEqual(names[2], long.upper()[:27] + ' (2)')
        # «Товары» уже занят основным листом — Excel не различает регистр.
        self.assertEqual(names[3], 'товары (2)')
        self.assertEqual(names[4], 'Без поставщика')
        self.assertTrue(all(len(name) <= 31 for name in names))


class ExportContentTests(TestCase):
    def setUp(self):
        # По цене первым идёт А, по остаткам — Б.
        self.a = Supplier.objects.create(name='А', price_priority=1, stock_priority=2)
        self.b = Supplier.objects.create(name='Б', price_priority=2, stock_priority=1)
        brand = Brand.objects.create(pim_id='br-1', name='Grohe')
        category = Category.objects.create(name='Сантехника')
        self.product = Product.objects.create(pim_id='p-1', number='SKU-1', name='Смеситель',
                                              brand=brand)
        self.product.categories.add(category)
        MainProduct.objects.create(
            product=self.product, supplier=self.a, article='A-1', name='Смеситель А',
            prime_cost=Decimal('0'), basic_price=Decimal('100'), m_price=None, stock=0)
        MainProduct.objects.create(
            product=self.product, supplier=self.b, article='B-1', name='Смеситель Б',
            prime_cost=Decimal('80'), basic_price=Decimal('120'), m_price=None, stock=None)

    def export(self, columns, query=''):
        content, count = ProductExporter(QueryDict(query), columns).build()
        return read_sheets(content), count

    def test_main_values_on_products_sheet_suppliers_on_own_sheets(self):
        sheets, count = self.export(['prime_cost', 'basic_price', 'm_price', 'stock'])
        self.assertEqual(count, 1)
        # Листы поставщиков — в порядке приоритета по цене.
        self.assertEqual(list(sheets), ['Товары', 'А', 'Б'])

        header, row = sheets['Товары']
        self.assertEqual(header, ['Артикул', 'Название', 'Бренд', 'Категории',
                                  'Себестоимость (основная)', 'Базовая цена (основная)',
                                  'Цена ИМ (основная)', 'Остаток (основной)'])
        values = dict(zip(header, row))
        self.assertEqual(values['Артикул'], 'SKU-1')
        self.assertEqual(values['Бренд'], 'Grohe')
        self.assertEqual(values['Категории'], 'Сантехника')
        # У А (приоритет по цене 1) себестоимость 0 — основная берётся у Б.
        self.assertEqual(values['Себестоимость (основная)'], 80)
        self.assertEqual(values['Базовая цена (основная)'], 100)
        # Цены ИМ нет ни у кого.
        self.assertIsNone(values['Цена ИМ (основная)'])
        # Остаток: у Б (приоритет по остаткам 1) NULL, у А 0 — основной 0, а не пусто.
        self.assertEqual(values['Остаток (основной)'], 0)

        self.assertEqual(sheets['А'][0], ['Артикул', 'Название', 'Себестоимость',
                                          'Базовая цена', 'Цена ИМ', 'Остаток'])
        self.assertEqual(sheets['А'][1], ['SKU-1', 'Смеситель', 0, 100, None, 0])
        self.assertEqual(sheets['Б'][1], ['SKU-1', 'Смеситель', 80, 120, None, None])

    def test_other_supplier_columns_only_on_supplier_sheets(self):
        sheets, _ = self.export(['article', 'supplier__msg_available', 'actions'])
        self.assertEqual(sheets['Товары'][0], ['Артикул', 'Название', 'Бренд', 'Категории'])
        self.assertEqual(sheets['А'][0], ['Артикул', 'Название', 'Артикул поставщика',
                                          'Поставщик • Сообщение при наличии'])
        self.assertEqual(sheets['А'][1][2], 'A-1')
        self.assertEqual(sheets['Б'][1][2], 'B-1')

    def test_supplier_sheet_holds_only_its_products_in_page_order(self):
        second = Product.objects.create(pim_id='p-2', number='SKU-0', name='Ванна')
        MainProduct.objects.create(product=second, supplier=self.b, article='B-2', name='Ванна')
        sheets, count = self.export(['article'], query='sort=number')
        self.assertEqual(count, 2)
        self.assertEqual([row['Артикул'] for row in as_dicts(sheets['Товары'])], ['SKU-0', 'SKU-1'])
        self.assertEqual([row['Артикул'] for row in as_dicts(sheets['А'])], ['SKU-1'])
        self.assertEqual([row['Артикул'] for row in as_dicts(sheets['Б'])], ['SKU-0', 'SKU-1'])

    def test_no_supplier_columns_no_supplier_sheets(self):
        sheets, _ = self.export(['ean'])
        self.assertEqual(list(sheets), ['Товары'])

    def test_every_selectable_column_is_writable(self):
        """«Поставщик» — объект модели, даты — с таймзоной: openpyxl не пишет ни то, ни другое."""
        MainProduct.objects.update(price_updated_at=timezone.now())
        sheets, _ = self.export(list(COLUMN_LABELS))
        values = as_dicts(sheets['А'])[0]
        self.assertEqual(values['Поставщик'], 'А')
        self.assertIsNotNone(values['Последнее обновление цены'])

    def test_several_rows_of_one_supplier_fall_through_too(self):
        MainProduct.objects.filter(supplier=self.b).delete()
        MainProduct.objects.create(product=self.product, supplier=self.a, article='A-2',
                                   name='Смеситель А2', m_price=Decimal('55'))
        sheets, _ = self.export(['m_price', 'article'])
        values = as_dicts(sheets['А'])[0]
        self.assertEqual(values['Цена ИМ'], 55)
        self.assertEqual(values['Артикул поставщика'], 'A-1; A-2')
        self.assertEqual(as_dicts(sheets['Товары'])[0]['Цена ИМ (основная)'], 55)

    def test_supplier_sheets_only_for_suppliers_of_exported_products(self):
        other = Product.objects.create(pim_id='p-2', number='SKU-2', name='Ванна')
        c = Supplier.objects.create(name='В', price_priority=3)
        MainProduct.objects.create(product=other, supplier=c, article='C-1', name='Ванна')
        sheets, count = self.export(['m_price'], query='search=SKU-1')
        self.assertEqual(count, 1)
        self.assertNotIn('В', sheets)


class ExportLevelTests(TestCase):
    """Цены: на уровне побеждает поставщик с минимальной себестоимостью.
    Остаток: максимальный на уровне."""

    def setUp(self):
        self.a = Supplier.objects.create(name='А', price_priority=1, stock_priority=1)
        self.b = Supplier.objects.create(name='Б', price_priority=1, stock_priority=1)
        self.c = Supplier.objects.create(name='В', price_priority=2, stock_priority=2)
        self.product = Product.objects.create(pim_id='p-1', number='SKU-1', name='Смеситель')

    def row(self, supplier, **values):
        return MainProduct.objects.create(product=self.product, supplier=supplier,
                                          name=self.product.name, **values)

    def main(self, columns):
        content, _ = ProductExporter(QueryDict(''), columns).build()
        return as_dicts(read_sheets(content)['Товары'])[0]

    def test_all_prices_from_winner_by_prime_cost_max_stock(self):
        self.row(self.a, prime_cost=Decimal('90'), basic_price=Decimal('120'), stock=2)
        self.row(self.b, prime_cost=Decimal('95'), basic_price=Decimal('100'), stock=5)
        # Нижний уровень дешевле и с большим остатком — но он не решает.
        self.row(self.c, prime_cost=Decimal('40'), basic_price=Decimal('50'), stock=99)
        values = self.main(['prime_cost', 'basic_price', 'stock'])
        # Победитель — А (себестоимость 90): и базовая цена его, хоть у Б она ниже.
        self.assertEqual(values['Себестоимость (основная)'], 90)
        self.assertEqual(values['Базовая цена (основная)'], 120)
        self.assertEqual(values['Остаток (основной)'], 5)

    def test_winner_chosen_even_when_prime_cost_not_exported(self):
        self.row(self.a, prime_cost=Decimal('90'), basic_price=Decimal('120'))
        self.row(self.b, prime_cost=Decimal('80'), basic_price=Decimal('130'))
        self.assertEqual(self.main(['basic_price'])['Базовая цена (основная)'], 130)

    def test_missing_price_of_winner_taken_from_next_in_level(self):
        self.row(self.a, prime_cost=Decimal('90'), basic_price=Decimal('0'), m_price=Decimal('150'))
        self.row(self.b, prime_cost=Decimal('95'), basic_price=Decimal('110'), m_price=Decimal('140'))
        self.row(self.c, prime_cost=Decimal('10'), basic_price=Decimal('20'))
        values = self.main(['basic_price', 'm_price'])
        self.assertEqual(values['Цена ИМ (основная)'], 150)
        # У победителя базовой цены нет — берётся у следующего на его уровне, не с нижнего.
        self.assertEqual(values['Базовая цена (основная)'], 110)

    def test_supplier_without_prime_cost_loses_to_one_with(self):
        self.row(self.a, basic_price=Decimal('50'))
        self.row(self.b, prime_cost=Decimal('95'), basic_price=Decimal('110'))
        self.assertEqual(self.main(['basic_price'])['Базовая цена (основная)'], 110)

    def test_level_without_values_falls_through(self):
        self.row(self.a, basic_price=Decimal('0'), stock=0)
        self.row(self.b, basic_price=None, stock=None)
        self.row(self.c, basic_price=Decimal('50'), stock=4)
        values = self.main(['basic_price', 'stock'])
        self.assertEqual(values['Базовая цена (основная)'], 50)
        self.assertEqual(values['Остаток (основной)'], 4)

    def test_unranked_are_one_level_below_ranked(self):
        alpha = Supplier.objects.create(name='Альфа')
        yashma = Supplier.objects.create(name='Яшма')
        self.row(alpha, prime_cost=Decimal('250'), basic_price=Decimal('300'), stock=1)
        self.row(yashma, prime_cost=Decimal('150'), basic_price=Decimal('200'), stock=8)
        values = self.main(['basic_price', 'stock'])
        # Не «Альфа» по алфавиту, а победитель по себестоимости и максимум остатка.
        self.assertEqual(values['Базовая цена (основная)'], 200)
        self.assertEqual(values['Остаток (основной)'], 8)
        self.row(self.c, basic_price=Decimal('900'), stock=1)
        values = self.main(['basic_price', 'stock'])
        self.assertEqual(values['Базовая цена (основная)'], 900)
        self.assertEqual(values['Остаток (основной)'], 1)

    def test_several_rows_of_one_supplier_cheapest_row_wins(self):
        self.row(self.a, prime_cost=Decimal('50'), m_price=Decimal('70'), stock=1)
        self.row(self.a, prime_cost=Decimal('40'), m_price=Decimal('75'), stock=3)
        content, _ = ProductExporter(QueryDict(''), ['m_price', 'stock']).build()
        sheet = as_dicts(read_sheets(content)['А'])[0]
        self.assertEqual(sheet['Цена ИМ'], 75)
        self.assertEqual(sheet['Остаток'], 3)

    def test_supplier_currency_prices_have_no_main_value(self):
        self.row(self.a, basic_price=Decimal('100'))
        columns = ['supplier_product_price', 'supplier_product_rrp',
                   'supplier_product_discount_price', 'basic_price']
        content, _ = ProductExporter(QueryDict(''), columns).build()
        sheets = read_sheets(content)
        header = sheets['Товары'][0]
        self.assertEqual(header[4:], ['Базовая цена (основная)'])
        # На листе поставщика они остаются.
        self.assertEqual(sheets['А'][0][2:], [COLUMN_LABELS[key] for key in columns])


class ExportOrderTests(TestCase):
    """Файл идёт в том же порядке, что и страница."""

    def setUp(self):
        self.user = User.objects.create_user(username='tester', password='pw')
        self.client.force_login(self.user)
        supplier = Supplier.objects.create(name='Поставщик')
        root = Category.objects.create(name='Б-категория')
        other = Category.objects.create(name='А-категория')
        for index, (number, category, stock) in enumerate([
            ('N-3', root, 5), ('N-1', other, 1), ('N-2', root, 9), ('N-4', None, 3),
        ]):
            product = Product.objects.create(pim_id=f'p-{index}', number=number,
                                             name=f'Кран {number}', raw_data={'name': f'Кран {number}'})
            if category:
                product.categories.add(category)
            product.rebuild_search_vector()
            MainProduct.objects.create(product=product, supplier=supplier, article=number,
                                       name=f'Кран {number}', stock=stock)

    def assert_same_order(self, query):
        response = self.client.get(reverse('products') + '?' + query)
        page = [row.record.pk for _, row in response.context['product_rows']]
        self.assertEqual(ordered_product_pks(QueryDict(query)), page)

    def test_default_by_category(self):
        self.assert_same_order('')

    def test_sorted_by_column(self):
        self.assert_same_order('sort=number')
        self.assert_same_order('sort=-total_stock')

    def test_search(self):
        self.assert_same_order('search=Кран')
        self.assert_same_order('search=Кран&sort=-number')


class ExportViewTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='tester', password='pw')
        self.client.force_login(self.user)

    def test_post_queues_task_with_page_query_without_page_number(self):
        with mock.patch('product.views.export_products_task.delay') as delay:
            response = self.client.post(reverse('product-export'),
                                        {'query': '?search=кран&sort=number&page=3'})
        self.assertEqual(response.status_code, 200)
        kwargs = delay.call_args.kwargs
        self.assertEqual(QueryDict(kwargs['query']).dict(), {'search': 'кран', 'sort': 'number'})
        self.assertEqual(kwargs['user_id'], self.user.pk)
        self.assertIn('stock', kwargs['columns'])

    def test_download_only_for_owner(self):
        export = ProductExport(user=self.user, rows_count=1)
        export.file.save('products-test.xlsx', ContentFile(b'data'), save=True)
        response = self.client.get(reverse('product-export-download', kwargs={'pk': export.pk}))
        self.assertEqual(response.status_code, 200)

        stranger = User.objects.create_user(username='other', password='pw')
        self.client.force_login(stranger)
        response = self.client.get(reverse('product-export-download', kwargs={'pk': export.pk}))
        self.assertEqual(response.status_code, 404)

    def test_task_notifies_with_download_link(self):
        from core.models import PersistentNotification
        from product.tasks import export_products_task

        Product.objects.create(pim_id='p-1', number='SKU-1', name='Смеситель')
        export_products_task(query='', columns=['stock'], user_id=self.user.pk)
        export = ProductExport.objects.get(user=self.user)
        self.assertEqual(export.rows_count, 1)
        notification = PersistentNotification.objects.get(user=self.user)
        self.assertEqual(notification.link,
                         reverse('product-export-download', kwargs={'pk': export.pk}))
        self.assertEqual(notification.kind, 'export')  # удаляется только вручную


def read_csv(content):
    """csv полного экспорта: (заголовок, строки словарями по заголовку)."""
    rows = list(csv.reader(StringIO(content.decode('utf-8-sig')), delimiter=';'))
    return rows[0], as_dicts(rows)


class FullCsvExportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(username='admin', password='pw')
        # По цене первым идёт А, по остаткам — Б.
        self.a = Supplier.objects.create(name='А', price_priority=1, stock_priority=2)
        self.b = Supplier.objects.create(name='Б', price_priority=2, stock_priority=1)
        self.product = Product.objects.create(pim_id='p-1', number='SKU-1', name='Смеситель')
        MainProduct.objects.create(
            product=self.product, supplier=self.a, article='A-1', name='Смеситель А',
            prime_cost=Decimal('0'), basic_price=Decimal('100'), stock=0)
        MainProduct.objects.create(
            product=self.product, supplier=self.b, article='B-1', name='Смеситель Б',
            prime_cost=Decimal('80'), basic_price=Decimal('120'), stock=7)

    def export(self):
        export = build_full_csv_export(self.user.pk)
        with export.file.open('rb') as file:
            return export, read_csv(file.read())

    def test_main_values_then_every_supplier_block_in_price_priority(self):
        Product.objects.create(pim_id='p-2', number='SKU-2', name='Без поставщиков')
        export, (header, rows) = self.export()
        self.assertEqual(export.rows_count, 2)
        self.assertTrue(export.file.name.endswith('.csv'))

        labels = [COLUMN_LABELS[key] for key in FULL_EXPORT_COLUMNS]
        self.assertEqual(header, (
            ['Артикул', 'Название', 'Бренд', 'Категории']
            + [f'{label} (основной)' if key == 'stock' else f'{label} (основная)'
               for key, label in zip(FULL_EXPORT_COLUMNS, labels)
               if key == 'stock' or key in MAIN_PRICE_COLUMNS]
            + [f'А • {label}' for label in labels]
            + [f'Б • {label}' for label in labels]))

        row = next(row for row in rows if row['Артикул'] == 'SKU-1')
        # У А (приоритет по цене 1) себестоимость 0 — основная берётся у Б.
        self.assertEqual(row['Себестоимость (основная)'], '80,00')
        self.assertEqual(row['Базовая цена (основная)'], '100,00')
        # По остаткам первым идёт Б.
        self.assertEqual(row['Остаток (основной)'], '7')
        self.assertEqual(row['Цена ИМ (основная)'], '')
        # Ноль остаётся нулём, а не пустой ячейкой (main_value).
        self.assertEqual(row['А • Себестоимость'], '0')
        self.assertEqual(row['Б • Базовая цена'], '120,00')
        self.assertEqual(row['А • Остаток'], '0')

        # Товар без строк поставщиков тоже в файле — с пустыми ценами.
        bare = next(row for row in rows if row['Артикул'] == 'SKU-2')
        self.assertEqual(bare['Базовая цена (основная)'], '')
        self.assertEqual(bare['А • Базовая цена'], '')

    def test_several_rows_of_one_supplier_fall_through(self):
        MainProduct.objects.filter(supplier=self.b).delete()
        MainProduct.objects.create(product=self.product, supplier=self.a, article='A-2',
                                   name='Смеситель А2', m_price=Decimal('55'))
        _, (header, rows) = self.export()
        self.assertFalse(any(title.startswith('Б • ') for title in header))
        self.assertEqual(rows[0]['А • Цена ИМ'], '55,00')
        self.assertEqual(rows[0]['Цена ИМ (основная)'], '55,00')

    def test_admin_button_posts_and_queues_task(self):
        self.client.force_login(self.user)
        changelist = reverse('admin:product_product_changelist')
        url = reverse('admin:product_product_export_full_csv')
        self.assertContains(self.client.get(changelist), url)

        self.assertEqual(self.client.get(url).status_code, 405)
        with mock.patch('product.admin.export_products_full_csv_task.delay') as delay:
            response = self.client.post(url)
        self.assertRedirects(response, changelist)
        delay.assert_called_once_with(user_id=self.user.pk)

    def test_admin_export_needs_view_permission(self):
        staff = User.objects.create_user(username='staff', password='pw', is_staff=True)
        self.client.force_login(staff)
        with mock.patch('product.admin.export_products_full_csv_task.delay') as delay:
            response = self.client.post(reverse('admin:product_product_export_full_csv'))
        self.assertEqual(response.status_code, 403)
        delay.assert_not_called()

    def test_task_notifies_and_download_is_csv(self):
        from core.models import PersistentNotification
        from product.tasks import export_products_full_csv_task

        export_products_full_csv_task(user_id=self.user.pk)
        export = ProductExport.objects.get(user=self.user)
        notification = PersistentNotification.objects.get(user=self.user)
        self.assertEqual(notification.link,
                         reverse('product-export-download', kwargs={'pk': export.pk}))

        self.client.force_login(self.user)
        response = self.client.get(notification.link)
        self.assertEqual(response.status_code, 200)
        self.assertIn('.csv', response['Content-Disposition'])
