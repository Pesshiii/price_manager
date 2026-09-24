"""Экспорт товаров: страница в xlsx и полный каталог в csv.

Полный csv (кнопка в админке товаров) — FullCsvExporter: те же правила, но
весь каталог, фиксированные колонки и поставщики колонками, а не листами.

xlsx — та же выборка, что на странице, тот же порядок, выбранные колонки.
Листы:

- «Товары» — строка на товар (Product), как строка страницы: колонки товара
  и основные цены и остаток — по уровням приоритета поставщиков
  (Supplier.price_priority / stock_priority), см. main_value_cells. Основное
  значение есть только у цен в тенге (MP_PRICES) и у остатка: цены
  SupplierProduct — в валюте поставщика, между поставщиками их не сравнить.
- по листу на поставщика, в порядке приоритета по цене — все выбранные
  колонки строк поставщиков (цены, остаток, артикул поставщика, статус
  наличия, …) с его значениями. На листе только товары, которые у него есть,
  в том же порядке, что на «Товарах»; артикул и название повторены, чтобы
  лист читался сам по себе.

Поставщики — все, у кого есть строка хотя бы у одного выгружаемого товара, а
не все поставщики базы: иначе фильтр по одному поставщику давал бы файл из
пустых листов остальных.
"""
import csv
import tempfile
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO, TextIOWrapper

from django.core.files import File
from django.core.files.base import ContentFile
from django.db.models import Prefetch
from django.http import QueryDict
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

from main_product_manager.models import MP_PRICES, MainProduct
from supplier_manager.models import Supplier

from .columns import COLUMN_LABELS, SUPPLIER_ROW_COLUMNS
from .filters import CATEGORY_LABEL_DEPTH, ProductFilter, search_terms
from .models import Category, Product, ProductExport
from .tables import ProductTable, SupplierRowTable, best_match_groups_first, category_path

# Параметр сортировки таблицы — тот же, что читает ProductPage.groups_by_category.
SORT_PARAM = ProductTable._meta.prefix + ProductTable._meta.order_by_field

# Цены SupplierProduct — аннотации views._latest_supplier_prices, не поля MainProduct.
SUPPLIER_PRODUCT_PRICES = ['supplier_product_price', 'supplier_product_rrp',
                           'supplier_product_discount_price']
PRICE_COLUMNS = set(MP_PRICES) | set(SUPPLIER_PRODUCT_PRICES)

# Колонки, у которых на «Товарах» есть основное значение: цены в тенге и
# остаток. Остальные колонки строк поставщиков — только на листах поставщиков.
MAIN_PRICE_COLUMNS = set(MP_PRICES)
STOCK_COLUMN = 'stock'

# Колонки товара, которые в файл не идут: «Действия» — кнопки, фото — картинка.
NOT_EXPORTED = {'actions', 'photo'}

NO_SUPPLIER = 'Без поставщика'

MAIN_SHEET = 'Товары'
PRODUCT_TITLES = ['Артикул', 'Название', 'Бренд', 'Категории']
# Начало строки на листе поставщика — по нему лист сверяется с «Товарами».
IDENTITY_TITLES = ['Артикул', 'Название']

# Полный csv-экспорт (FullCsvExporter): все цены и остаток, в порядке выбора колонок.
FULL_EXPORT_COLUMNS = [key for key in SUPPLIER_ROW_COLUMNS
                       if key in PRICE_COLUMNS or key == STOCK_COLUMN]

# Имя листа Excel: не длиннее 31 символа, без []:*?/\ и не пустое.
SHEET_NAME_LIMIT = 31
SHEET_NAME_FORBIDDEN = str.maketrans({char: ' ' for char in '[]:*?/\\'})


def sheet_names(names, taken=(MAIN_SHEET,)) -> list[str]:
    """Имена листов для поставщиков: допустимые в Excel и попарно разные.

    Excel сравнивает имена листов без учёта регистра, и обрезка до 31 символа
    может свести два длинных имени к одному — такие получают суффикс « (2)».
    """
    used = {name.lower() for name in taken}
    result = []
    for name in names:
        base = (name.translate(SHEET_NAME_FORBIDDEN).strip().strip("'")
                or NO_SUPPLIER)[:SHEET_NAME_LIMIT]
        candidate, number = base, 2
        while candidate.lower() in used:
            suffix = f' ({number})'
            candidate = base[:SHEET_NAME_LIMIT - len(suffix)] + suffix
            number += 1
        used.add(candidate.lower())
        result.append(candidate)
    return result

CHUNK_SIZE = 1000


def ordered_product_pks(params) -> list[int]:
    """pk товаров в том порядке, в каком их показывает ProductPage.

    Повторяет цепочку страницы: фильтр → выдача по категориям или сортировка
    по колонке → при поиске в выдаче по категориям — best_match_groups_first
    → order_by таблицы. Разойдись они, файл лёг бы не в том порядке, что
    экран; это держит тест test_export.ExportOrderTests.

    Невалидный фильтр — пустая выборка, как у FilterView со strict.
    """
    from .views import _base_queryset

    sort = params.get(SORT_PARAM)
    by_category = not sort
    filterset = ProductFilter(params, queryset=_base_queryset(by_category=by_category))
    if not filterset.is_valid():
        return []
    queryset = filterset.qs
    if by_category and search_terms(params.get('search')):
        queryset = best_match_groups_first(queryset)
    table = ProductTable(queryset, order_by=sort or None)
    return list(table.data.data.values_list('pk', flat=True))


def export_columns(selected) -> tuple[list[str], list[str]]:
    """Выбранные колонки, разделённые на колонки товара и строк поставщиков."""
    product_columns = [key for key in selected
                       if key not in SUPPLIER_ROW_COLUMNS and key not in NOT_EXPORTED]
    supplier_columns = [key for key in selected
                        if key in SUPPLIER_ROW_COLUMNS and key not in NOT_EXPORTED]
    return product_columns, supplier_columns


def is_zero(value) -> bool:
    return value is None or value == 0


def main_value(levels, pick):
    """Основное значение по уровням: pick (min или max) из ненулевых значений
    первого уровня, где они есть.

    levels — значения, сгруппированные по уровням, сверху вниз; уровень из
    одного значения — просто «первое ненулевое по порядку». Ноль и пусто
    пропускаются — у самого приоритетного поставщика ноль чаще значит «цены
    нет», чем «бесплатно», и тогда решает следующий уровень. Если ненулевого
    нет ни на одном уровне: 0, если хоть у кого-то 0 (остаток
    синхронизировался, товара нет), иначе пусто (ни у кого нет данных) — NULL
    и 0 в остатке разные вещи.
    """
    seen = False
    for level in levels:
        level = list(level)
        non_zero = [value for value in level if not is_zero(value)]
        if non_zero:
            return pick(non_zero)
        seen = seen or any(value is not None for value in level)
    return 0 if seen else None


def cost_key(cost):
    """Ключ сортировки по себестоимости: меньшая — раньше, ноль и пусто — в конце."""
    return (is_zero(cost), cost or 0)


def winner_order(level, costs) -> list:
    """Поставщики уровня по цене: победитель (минимальная себестоимость) первым.

    Без себестоимости — в конце. sorted устойчив, так что при равной
    себестоимости остаётся порядок уровня — по имени.
    """
    return sorted(level, key=lambda pk: cost_key(costs.get(pk)))


def ranked_suppliers(supplier_ids, priority_field) -> list:
    """Поставщики по приоритету: меньше — выше, непроранжированные в конце.

    Внутри уровня — по имени. None в supplier_ids — строки без поставщика;
    они всегда последние.
    """
    suppliers = list(Supplier.objects.filter(pk__in=[pk for pk in supplier_ids if pk]))
    suppliers.sort(key=lambda s: (getattr(s, priority_field) is None,
                                  getattr(s, priority_field) or 0, s.name, s.pk))
    result = [(supplier.pk, supplier.name) for supplier in suppliers]
    if None in supplier_ids:
        result.append((None, NO_SUPPLIER))
    return result


def supplier_levels(supplier_ids, priority_field) -> list[list]:
    """pk поставщиков, сгруппированные по уровням приоритета, сверху вниз.

    Один номер — один уровень. Все непроранжированные — один общий уровень
    ниже проранжированных: порядок между ними ничего не значит. Строки без
    поставщика (None) — отдельный уровень в самом конце. Внутри уровня — по имени.
    """
    levels = defaultdict(list)
    unranked = []
    for pk, level in (Supplier.objects.filter(pk__in=[pk for pk in supplier_ids if pk])
                      .order_by('name', 'pk').values_list('pk', priority_field)):
        (unranked if level is None else levels[level]).append(pk)
    result = [levels[level] for level in sorted(levels)]
    if unranked:
        result.append(unranked)
    if None in supplier_ids:
        result.append([None])
    return result


def _chunks(items, size=CHUNK_SIZE):
    for start in range(0, len(items), size):
        yield items[start:start + size]


def _excel_value(value):
    """Значение, которое openpyxl примет.

    datetime с таймзоной он не пишет — переводим в местное время. Объекты
    моделей (колонка «Поставщик» — это Supplier) — строкой, как на экране.
    """
    if isinstance(value, datetime):
        return timezone.localtime(value).replace(tzinfo=None) if value.tzinfo else value
    if value is None or isinstance(value, (str, int, float, Decimal, date)):
        return value
    return str(value)


def _csv_value(value):
    """Ячейка csv: пусто — пустая строка, дробные — с запятой.

    Файл для Excel с русской локалью: с точкой «3.10» он прочтёт как 3 октября.
    """
    if value is None:
        return ''
    if isinstance(value, (Decimal, float)):
        return str(value).replace('.', ',')
    return _excel_value(value)


def _joined(values):
    """Несколько строк одного поставщика: различающиеся значения через «; »."""
    distinct = []
    for value in values:
        if value in (None, '') or value in distinct:
            continue
        distinct.append(value)
    if not distinct:
        return None
    if len(distinct) == 1:
        return distinct[0]
    return '; '.join(str(value) for value in distinct)


class ProductExporter:
    """Собирает xlsx по параметрам страницы и выбранным колонкам."""

    def __init__(self, params, selected_columns):
        self.params = params
        self.product_columns, self.supplier_columns = export_columns(selected_columns)
        # Та же отрисовка «Статуса наличия» и «Срока поставки», что на экране.
        self.row_table = SupplierRowTable([])

    def main_products(self, pks):
        queryset = (MainProduct.objects.filter(product_id__in=pks)
                    .select_related('supplier', 'supplier__currency').order_by('pk'))
        if set(self.supplier_columns) & set(SUPPLIER_PRODUCT_PRICES):
            from .views import _latest_supplier_prices
            queryset = queryset.annotate(**_latest_supplier_prices())
        return queryset

    def products(self, pks):
        categories = Category.objects.select_related(CATEGORY_LABEL_DEPTH)
        queryset = (Product.objects.filter(pk__in=pks).select_related('brand')
                    .prefetch_related(Prefetch('categories', queryset=categories))
                    .defer('search_vector'))
        return {product.pk: product for product in queryset}

    def cell(self, main_product, key):
        if key == 'stock_msg':
            return self.row_table.render_stock_msg(main_product)
        if key == 'delivery_days':
            return self.row_table.render_delivery_days(main_product)
        value = main_product
        for part in key.split('__'):
            value = getattr(value, part, None)
            if value is None:
                return None
        return value

    def main_titles(self) -> list[str]:
        """«Товары»: товар и основные цены и остаток (по уровням поставщиков)."""
        titles = PRODUCT_TITLES + [COLUMN_LABELS[key] for key in self.product_columns]
        for key in self.supplier_columns:
            if key in MAIN_PRICE_COLUMNS:
                titles.append(f'{COLUMN_LABELS[key]} (основная)')
            elif key == STOCK_COLUMN:
                titles.append(f'{COLUMN_LABELS[key]} (основной)')
        return titles

    def supplier_titles(self) -> list[str]:
        """Лист поставщика: артикул и название товара, затем все выбранные колонки."""
        return IDENTITY_TITLES + [COLUMN_LABELS[key] for key in self.supplier_columns]

    def product_cells(self, product, main_products):
        number = product.number or min(
            (mp.sku for mp in main_products if mp.sku), default=None)
        name = product.name or next((mp.name for mp in main_products if mp.name), '')
        categories = sorted(product.categories.all(), key=lambda c: (c.tree_id, c.lft))
        raw = product.raw_data or {}
        cells = [
            number,
            name,
            product.brand.name if product.brand else None,
            '; '.join(' / '.join(category_path(category)) for category in categories) or None,
        ]
        for key in self.product_columns:
            if key == 'tags':
                cells.append(', '.join(raw.get('tag') or []) or None)
            elif key == 'ean':
                cells.append(raw.get('ean') or None)
            elif key == 'pim_status':
                cells.append(raw.get('status') or None)
        return cells

    def supplier_values(self, main_products) -> dict:
        """{pk поставщика: {колонка: значение}} по строкам товара.

        Уникальности (товар, поставщик) у MainProduct нет, и строк у поставщика
        бывает несколько. Они сводятся тем же правилом, что поставщики одного
        уровня: цены — первое ненулевое по строкам от минимальной
        себестоимости, остаток — максимальный; текст — различающиеся значения
        через «; ».
        """
        by_supplier = defaultdict(list)
        for main_product in main_products:
            by_supplier[main_product.supplier_id].append(main_product)
        values = {}
        for supplier_pk, rows in by_supplier.items():
            rows = sorted(rows, key=lambda mp: cost_key(mp.prime_cost))
            values[supplier_pk] = {
                key: (main_value([[self.cell(mp, key) for mp in rows]], max)
                      if key == STOCK_COLUMN
                      else main_value([[self.cell(mp, key)] for mp in rows], min)
                      if key in PRICE_COLUMNS
                      else _joined(self.cell(mp, key) for mp in rows))
                for key in self.supplier_columns
            }
        return values

    @staticmethod
    def supplier_costs(main_products) -> dict:
        """{pk поставщика: себестоимость} — минимальная ненулевая по его строкам.

        Считается всегда, даже если колонка себестоимости не выбрана: по ней
        выбирается победитель уровня.
        """
        by_supplier = defaultdict(list)
        for main_product in main_products:
            by_supplier[main_product.supplier_id].append(main_product.prime_cost)
        return {pk: main_value([costs], min) for pk, costs in by_supplier.items()}

    def main_value_cells(self, values, costs, price_levels, stock_levels):
        """Основные цены и остаток по уровням поставщиков (main_value).

        Цены: на уровне побеждает поставщик с минимальной себестоимостью, и
        все цены берутся у него — чтобы цены строки были от одного поставщика.
        Нет у него какой-то цены (0 или пусто) — она берётся у следующего по
        себестоимости на том же уровне, затем с уровней ниже.

        Остаток: максимальный ненулевой на первом уровне по остаткам, где он есть.
        """
        price_order = [[pk] for level in price_levels
                       for pk in winner_order([pk for pk in level if pk in values], costs)]
        cells = []
        for key in self.supplier_columns:
            if key in MAIN_PRICE_COLUMNS:
                cells.append(main_value(([values[pk][key] for pk in single]
                                         for single in price_order), min))
            elif key == STOCK_COLUMN:
                cells.append(main_value(
                    ([values[pk][key] for pk in level if pk in values] for level in stock_levels),
                    max))
        return cells

    def suppliers(self, pks):
        """Поставщики выгружаемых товаров: [(pk, имя)] по приоритету цены (порядок
        листов и блоков) и уровни по цене и по остаткам (supplier_levels)."""
        supplier_ids = set()
        if self.supplier_columns:
            for chunk in _chunks(pks):
                supplier_ids.update(
                    MainProduct.objects.filter(product_id__in=chunk)
                    .order_by().values_list('supplier_id', flat=True).distinct())
        return (ranked_suppliers(supplier_ids, 'price_priority'),
                supplier_levels(supplier_ids, 'price_priority'),
                supplier_levels(supplier_ids, 'stock_priority'))

    def rows(self, pks, price_levels, stock_levels):
        """По строке на товар в порядке pks: (ячейки товара, основные значения,
        {pk поставщика: {колонка: значение}} — только поставщики, у которых
        товар есть).

        Общая часть xlsx и csv: правила основного значения и свёртки строк
        одного поставщика у них одни.
        """
        for chunk in _chunks(pks):
            products = self.products(chunk)
            main_products = defaultdict(list)
            for main_product in self.main_products(chunk):
                main_products[main_product.product_id].append(main_product)
            # pk__in порядок не хранит — порядок страницы восстанавливается по chunk.
            for pk in chunk:
                product = products.get(pk)
                if product is None:
                    continue
                rows = main_products.get(pk, [])
                values = self.supplier_values(rows) if self.supplier_columns else {}
                yield (self.product_cells(product, rows),
                       self.main_value_cells(values, self.supplier_costs(rows),
                                             price_levels, stock_levels),
                       values)

    def build(self) -> tuple[bytes, int]:
        pks = ordered_product_pks(self.params)
        price_suppliers, price_levels, stock_levels = self.suppliers(pks)
        price_order = [pk for pk, _ in price_suppliers]

        workbook = Workbook(write_only=True)
        bold = Font(bold=True)

        def add_sheet(title, header):
            sheet = workbook.create_sheet(title)
            sheet.freeze_panes = 'A2'
            cells = []
            for value in header:
                cell = WriteOnlyCell(sheet, value=value)
                cell.font = bold
                cells.append(cell)
            sheet.append(cells)
            return sheet

        main_sheet = add_sheet(MAIN_SHEET, self.main_titles())
        # Листы поставщиков — в порядке приоритета по цене, как колонки выбора.
        supplier_sheets = {
            pk: add_sheet(title, self.supplier_titles())
            for (pk, _), title in zip(price_suppliers,
                                      sheet_names(name for _, name in price_suppliers))
        }

        for product_cells, main_cells, values in self.rows(pks, price_levels, stock_levels):
            main_sheet.append([_excel_value(value) for value in product_cells + main_cells])
            identity = product_cells[:len(IDENTITY_TITLES)]
            for supplier_pk in price_order:
                if supplier_pk in values:
                    supplier_sheets[supplier_pk].append([_excel_value(value) for value in (
                        identity + [values[supplier_pk][key] for key in self.supplier_columns])])

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue(), len(pks)


class FullCsvExporter(ProductExporter):
    """Полный экспорт каталога в csv — кнопка в админке товаров.

    Весь каталог (фильтров нет, порядок — страница по умолчанию), одна
    строка на товар: товар, основные цены и остаток по уровням
    поставщиков, затем по блоку на каждого поставщика — все его цены и
    остаток, колонки «<поставщик> • <колонка>», поставщики по приоритету
    цены. У csv один лист, поэтому поставщики — колонками, а не листами.

    Разделитель «;», BOM и дробные с запятой — так файл сразу открывает Excel
    с русской локалью, числа остаются числами.
    """

    DELIMITER = ';'
    ENCODING = 'utf-8-sig'

    def __init__(self):
        super().__init__(QueryDict(), FULL_EXPORT_COLUMNS)

    def write(self, stream) -> int:
        pks = ordered_product_pks(self.params)
        price_suppliers, price_levels, stock_levels = self.suppliers(pks)
        price_order = [pk for pk, _ in price_suppliers]
        labels = [COLUMN_LABELS[key] for key in self.supplier_columns]

        writer = csv.writer(stream, delimiter=self.DELIMITER)
        writer.writerow(self.main_titles() + [
            f'{name} • {label}' for _, name in price_suppliers for label in labels])
        empty = dict.fromkeys(self.supplier_columns)
        for product_cells, main_cells, values in self.rows(pks, price_levels, stock_levels):
            supplier_cells = [values.get(pk, empty)[key]
                              for pk in price_order for key in self.supplier_columns]
            writer.writerow([_csv_value(value)
                             for value in product_cells + main_cells + supplier_cells])
        return len(pks)


def build_product_export(params, selected_columns, user_id) -> ProductExport:
    content, rows_count = ProductExporter(params, selected_columns).build()
    export = ProductExport(user_id=user_id, rows_count=rows_count)
    # Имя на диске — ASCII; человеческое имя подставляет представление скачивания.
    export.file.save(f'products-{user_id}-{timezone.now():%Y%m%d-%H%M%S}.xlsx',
                     ContentFile(content), save=True)
    return export


def build_full_csv_export(user_id) -> ProductExport:
    # Весь каталог — сотни колонок на ~158 тыс. строк: пишется во временный
    # файл, а не собирается в памяти.
    with tempfile.TemporaryFile() as raw:
        stream = TextIOWrapper(raw, encoding=FullCsvExporter.ENCODING, newline='')
        rows_count = FullCsvExporter().write(stream)
        stream.flush()
        raw.seek(0)
        export = ProductExport(user_id=user_id, rows_count=rows_count)
        export.file.save(f'products-full-{user_id}-{timezone.now():%Y%m%d-%H%M%S}.csv',
                         File(raw), save=True)
        stream.detach()
    return export
